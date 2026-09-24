"""Derive the twelve-project Word template from the retained redesigned template."""

from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "assets/templates/planificacion-anual-inicial-redisenada.docx"
TARGET = ROOT / "assets/templates/planificacion-anual-inicial-flexible.docx"


def text_of(element):
    return "".join(node.text or "" for node in element.xpath(".//w:t"))


def set_text(element, value):
    nodes = element.xpath(".//w:t")
    if not nodes:
        raise ValueError(f"Element without text slot: {value}")
    nodes[0].text = value
    for node in nodes[1:]:
        node.text = ""


document = Document(SOURCE)
body = document.element.body
original = list(body)

# The source groups each project in ten body elements. Remove P13–P20 and
# retain its evaluation and closing sections.
assert "{{PROYECTO_13_TITULO}}" in text_of(original[163])
assert "VII. EVALUACIÓN" in text_of(original[241])
for element in original[161:241]:
    body.remove(element)

# The source splits the chronology across two mostly empty pages. Twelve rows
# fit in one table, so append P11 and P12 and remove the continuation page.
first_chronology = original[35]
second_chronology = original[39]
for row in second_chronology.xpath("./w:tr")[1:3]:
    first_chronology.append(deepcopy(row))
for element in [original[30], *original[36:40]]:
    body.remove(element)

# The two source callouts repeat what the section titles already say. Removing
# them gives the twelve-row chronology enough room to stay on one page.
for element in [original[27], original[28], original[33], original[34]]:
    body.remove(element)

# Reuse the template's heading and two-column data table for the independent
# initial stage. It has no project code or required final product.
heading = deepcopy(original[25])
set_text(heading, "IV. ACOGIDA, ADAPTACIÓN Y EVALUACIÓN DIAGNÓSTICA")
stage_table = deepcopy(original[11])
stage_rows = stage_table.xpath("./w:tr")
stage_fields = [
    ("Duración referencial", "{{ETAPA_DURACION}}"),
    ("Propósito", "{{ETAPA_PROPOSITO}}"),
    ("Qué hará la docente", "{{ETAPA_EXPERIENCIAS}}"),
    ("Qué observará", "{{ETAPA_OBSERVAR}}"),
    ("Con las familias", "{{ETAPA_FAMILIAS}}"),
    ("Foco diagnóstico", "{{ETAPA_FOCO}}"),
    ("Fechas previstas", "{{ETAPA_FECHAS}}"),
    ("Para tener en cuenta", "La adaptación puede continuar de manera diferenciada según las necesidades de cada niño."),
]
for row, (label, value) in zip(stage_rows, stage_fields, strict=True):
    cells = row.xpath("./w:tc")
    set_text(cells[0], label)
    set_text(cells[1], value)
at = body.index(original[25])
for item in (heading, deepcopy(original[26]), stage_table, deepcopy(original[26])):
    body.insert(at, item)
    at += 1
body.insert(body.index(original[25]), deepcopy(original[30]))

# Do not strand the next section heading at the foot of the stage page.
for paragraph in original[25].xpath(".//w:p"):
    properties = paragraph.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr")
    if properties is None:
        properties = OxmlElement("w:pPr")
        paragraph.insert(0, properties)
    properties.append(OxmlElement("w:keepNext"))

# Make the yearly summary about teaching periods. Keep the original table
# geometry, color, typography, and only four data rows.
for table in document.tables:
    if table.cell(0, 0).text.strip() == "Mes":
        rows = table.rows
        for row in list(rows)[5:]:
            table._tbl.remove(row._tr)
        set_text(table.rows[0].cells[0]._tc, "Periodo lectivo")
        set_text(table.rows[0].cells[1]._tc, "Fechas")
        set_text(table.rows[0].cells[2]._tc, "Proyectos previstos")
        for index, row in enumerate(list(table.rows)[1:], start=1):
            set_text(row.cells[0]._tc, f"Periodo {index}")
            set_text(row.cells[1]._tc, f"{{{{BLOQUE_{index}_FECHAS}}}}")
            set_text(row.cells[2]._tc, f"{{{{BLOQUE_{index}_CODIGOS}}}}")
    if table.cell(0, 0).text.strip() == "Cod.":
        for row in list(table.rows)[1:]:
            if any(f"PROYECTO_{index:02d}_" in cell.text for index in range(13, 21) for cell in row.cells):
                table._tbl.remove(row._tr)
for table in document.tables:
    content = text_of(table._tbl)
    if "{{PROYECTO_" in content and "_TITULO}}" in content:
        import re
        match = re.search(r"PROYECTO_(\d{2})_TITULO", content)
        if match and len(table.rows) == 1 and len(table.rows[0].cells) >= 2:
            index = int(match.group(1))
            if index <= 12:
                set_text(table.rows[0].cells[0]._tc, f"PERIODO {((index - 1) // 3) + 1}")

for table in document.tables:
    head = table.cell(0, 0).text.strip()
    if head.startswith("IV. DISTRIBUCIÓN DE PROYECTOS POR MES"):
        set_text(table._tbl, "V. ORGANIZACIÓN DE PROYECTOS POR PERIODO LECTIVO")
    elif head.startswith("V. CRONOGRAMA GENERAL DE PROYECTOS"):
        set_text(table._tbl, head.replace("V.", "VI.", 1))
    elif head.startswith("VI. DESARROLLO MENSUAL DE LA PLANIFICACIÓN ANUAL"):
        set_text(table._tbl, "VII. DESARROLLO DE LOS PROYECTOS")
    elif head.startswith("VII. EVALUACIÓN Y SEGUIMIENTO"):
        set_text(table._tbl, "VIII. EVALUACIÓN Y SEGUIMIENTO")
    elif head.startswith("VIII. OBSERVACIONES FINALES"):
        set_text(table._tbl, "IX. OBSERVACIONES FINALES")
    elif head.startswith("La cantidad de proyectos puede variar"):
        set_text(table._tbl, "Estas doce propuestas iniciales se organizan por periodos lectivos. La docente puede moverlas, sustituirlas o ajustarlas según lo que observe en el grupo.")
    elif head.startswith("Registra una fila por proyecto"):
        set_text(table._tbl, "Las fechas son tentativas. Cada proyecto empieza un lunes y termina un viernes dentro de semanas lectivas de la institución.")


def minimum_cell_margin(cell, side, twips):
    tc_pr = cell._tc.get_or_add_tcPr()
    current = tc_pr.xpath("./w:tcMar")
    tc_mar = current[0] if current else OxmlElement("w:tcMar")
    if not current:
        tc_pr.append(tc_mar)
    names = (side, "left" if side == "start" else "right") if side in ("start", "end") else (side,)
    existing = next((child for child in tc_mar if child.tag in {qn(f"w:{name}") for name in names}), None)
    if existing is None:
        existing = OxmlElement(f"w:{side}")
        tc_mar.append(existing)
    value = int(existing.get(qn("w:w"), "0"))
    if value < twips:
        existing.set(qn("w:w"), str(twips))
        existing.set(qn("w:type"), "dxa")


# Leave the two long overview tables at their designed size so the chronology
# stays on one page. Give the stage and twelve narrative cards more breathing
# room; each of those blocks has its own page space.
for table in document.tables:
    content = text_of(table._tbl)
    if "{{ETAPA_DURACION}}" in content:
        vertical = 105
    elif "_SITUACION}}" in content and "{{PROYECTO_" in content:
        vertical = 115
    else:
        continue
    for row in table.rows:
        for cell in row.cells:
            for side in ("top", "bottom"):
                minimum_cell_margin(cell, side, vertical)
            for paragraph in cell.paragraphs:
                if paragraph.text.strip():
                    paragraph.paragraph_format.line_spacing = 1.10

# Two adjacent tables touch in Word even when their inner cells have padding.
# Add a small external gap between the period summary and the chronology title.
for element in list(body):
    if element.tag == qn("w:tbl") and "Periodo lectivoFechasProyectos previstos" in text_of(element):
        gap = document.add_paragraph()
        gap.paragraph_format.space_before = Pt(0)
        gap.paragraph_format.space_after = Pt(0)
        gap.paragraph_format.line_spacing = Pt(12)
        gap.add_run(" ").font.size = Pt(2)
        element.addnext(gap._element)
        break

TARGET.parent.mkdir(parents=True, exist_ok=True)
document.save(TARGET)
assert "PROYECTO_13_" not in "".join(text_of(item) for item in document.element.body)
print(TARGET)
