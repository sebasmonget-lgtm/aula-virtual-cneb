"""Adapt the supplied diagnostic Word design to Ayni's confirmed data model.

Run once with the teacher's source DOCX. The generated asset is used at runtime;
the source file is never needed by the server.
"""

from __future__ import annotations

import argparse
import base64
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def set_cell(cell, label, marker):
    cell.text = f"{label}\n{{{{{marker}}}}}"
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = paragraph == cell.paragraphs[0]


def insert_gap_after(doc, element, points):
    gap = doc.add_paragraph()
    gap.paragraph_format.space_before = Pt(0)
    gap.paragraph_format.space_after = Pt(0)
    gap.paragraph_format.line_spacing = Pt(points)
    gap.add_run(" ").font.size = Pt(2)
    element.addnext(gap._element)
    return gap._element


def build(source: Path, target: Path):
    doc = Document(source)
    body = doc.element.body
    original = list(body)
    original_tables = list(doc.tables)

    # Keep the designed cover and sections I-V, VII-IX. The source template's
    # competence-by-competence interpretations and named follow-ups are not
    # canonical Ayni records and cannot be filled by inference.
    for index in [*range(30, 40), 43, 44, 45, 58]:
        body.remove(original[index])

    heading = original[28]
    nodes = heading.xpath('.//w:t')
    for index, node in enumerate(nodes):
        node.text = "VI. OBSERVACIONES POR COMPETENCIA" if index == 0 else ""
    instruction = original[29]
    nodes = instruction.xpath('.//w:t')
    for index, node in enumerate(nodes):
        node.text = (
            "Esta tabla muestra cuántos registros hay por competencia. "
            "No asigna niveles de logro ni reemplaza la interpretación de la docente."
        ) if index == 0 else ""

    table = doc.add_table(rows=2, cols=3)
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, (Inches(3.9), Inches(1.35), Inches(1.35))):
            cell.width = width
    for cell, label in zip(table.rows[0].cells, ["Competencia", "Niños observados", "Registros"]):
        cell.text = label
    for cell, marker in zip(table.rows[1].cells, ["COMPETENCY_NAME", "COMPETENCY_STUDENTS", "COMPETENCY_RECORDS"]):
        cell.text = "{{" + marker + "}}"
    for index, row in enumerate(table.rows):
        for cell in row.cells:
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "244667" if index == 0 else "EAF4FA")
            cell._tc.get_or_add_tcPr().append(shade)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(7)
                paragraph.paragraph_format.space_after = Pt(7)
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(9)
                    run.font.bold = index == 0
                    run.font.color.rgb = RGBColor(255, 255, 255) if index == 0 else RGBColor(27, 50, 75)
    body.remove(table._element)
    before_table = insert_gap_after(doc, instruction, 10)
    before_table.addnext(table._element)
    insert_gap_after(doc, table._element, 16)

    # The application stores one teacher-confirmed group synthesis. There are
    # no confirmed group fields for interests, barriers, or family conclusions.
    synthesis = original_tables[14]
    for row in list(synthesis.rows)[1:]:
        synthesis._element.remove(row._element)

    for table, row_index in [(original_tables[1], 4), (original_tables[3], 6)]:
        table.rows[row_index].cells[1].text = "{{PERIODO_DIAGNOSTICO}}"

    sources = original_tables[9]
    for row in [sources.rows[4], sources.rows[1]]:
        sources._element.remove(row._element)
    sources.rows[2].cells[1].text = "Registros de Ayni Aula"
    comment = sources.add_row()
    for cell, value in zip(comment.cells, [
        "Comentario de la docente", "Síntesis individual confirmada",
        "Reúne la interpretación de la docente sobre cada niño.",
        "{{ESTADO_COMENTARIOS}}",
    ]):
        cell.text = value
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), "F0F7FA")
        cell._tc.get_or_add_tcPr().append(shade)

    coverage_heading = original_tables[10]
    nodes = coverage_heading._element.xpath('.//w:t')
    for index, node in enumerate(nodes):
        node.text = "IV. COBERTURA DE REGISTROS AL DESCARGAR" if index == 0 else ""
    set_cell(original_tables[11].rows[1].cells[0], "Registros de observación", "N_EVIDENCIAS_REVISADAS")

    decisions = original_tables[24]
    set_cell(decisions.rows[0].cells[0], "Decisiones de la docente", "PRIORIDADES_DIAGNOSTICAS")
    set_cell(decisions.rows[0].cells[1], "Fortalezas para seguir aprovechando", "DIAGNOSTICO_FORTALEZAS")
    decisions._element.remove(decisions.rows[1]._element)

    references = original_tables[31]
    for row in list(references.rows)[2:]:
        references._element.remove(row._element)

    # A real logo replaces this tiny embedded image during export. Its presence
    # gives OOXML a stable drawing and relationship without editing the layout.
    logo_cell = original_tables[0].cell(0, 0)
    logo_cell.text = ""
    transparent_png = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/lYQAAAAASUVORK5CYII="
    )
    logo_cell.paragraphs[0].add_run().add_picture(BytesIO(transparent_png), width=Inches(0.70))

    target.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "Evaluación diagnóstica del aula"
    doc.core_properties.subject = "Informe pedagógico de Educación Inicial"
    doc.core_properties.comments = "Informe de Ayni Aula completado desde registros autorizados."
    doc.core_properties.author = "Ayni Aula"
    doc.core_properties.last_modified_by = "Ayni Aula"
    doc.save(target)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("target", type=Path)
    args = parser.parse_args()
    build(args.source, args.target)
